import csv
import io
import re
from datetime import datetime
from typing import Any, List, Optional

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from issues.models import Issue
from projects.models import Project, ProjectDoc, ProjectIntelligenceSummary, ProjectMembership, ProjectSignal


def sanitize_formula_injection(value: Any) -> Any:
    """
    Prevents CSV / Excel Formula Injection (CWE-1236).
    If a string begins with dangerous prefix characters (=, +, -, @),
    it is prepended with a single quote to force Excel/Sheets to treat it as plain text.
    """
    if value is None:
        return ""
    val_str = str(value)
    if val_str and val_str[0] in ("=", "+", "-", "@", "\t", "\r"):
        return f"'{val_str}"
    return val_str


def _header_font(color: str = "FFFFFF", size: int = 11, bold: bool = True) -> Font:
    return Font(name="Segoe UI", size=size, bold=bold, color=color)


def _cell_font(size: int = 10, bold: bool = False, color: str = "172B4D") -> Font:
    return Font(name="Segoe UI", size=size, bold=bold, color=color)


def _header_fill(hex_color: str = "0747A6") -> PatternFill:
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


def _zebra_fill(hex_color: str = "F4F5F7") -> PatternFill:
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


def _thin_border() -> Border:
    thin = Side(border_style="thin", color="DFE1E6")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _auto_fit_columns(ws, min_width: int = 12, max_width: int = 60):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val = cell.value
            if val is not None:
                lines = str(val).split("\n")
                max_len = max(max_len, max(len(line) for line in lines))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, min_width), max_width)


def generate_canonical_project_workbook(
    project: Project,
    issues_qs: Optional[Any] = None,
    user: Optional[Any] = None,
) -> io.BytesIO:
    """
    Generates the canonical 5-sheet NEXO PMO Excel Workbook:
    1. Project Overview
    2. Tasks
    3. Milestones
    4. Risks & Issues
    5. Lists & Definitions
    """
    wb = Workbook()
    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    border = _thin_border()
    navy_fill = _header_fill("0747A6")
    sub_fill = _header_fill("0052CC")
    accent_fill = _header_fill("172B4D")
    zebra = _zebra_fill("F8F9FA")

    # Fetch summary and intelligence
    summary = getattr(project, "intelligence_summary", None)
    norm_state = summary.normalized_state if summary else {}
    health_status = summary.health_status if summary else "ON_TRACK"
    health_rationale = summary.health_rationale if summary else "Project progressing in accordance with baseline."
    current_phase = summary.current_phase if summary else "Active Execution"
    executive_bullets = summary.executive_summary if summary else ""

    if issues_qs is None:
        issues_qs = project.issues.all().select_related("assignee", "reporter", "sprint")

    # =========================================================================
    # SHEET 1: PROJECT OVERVIEW
    # =========================================================================
    ws_overview = wb.create_sheet(title="Project Overview")
    ws_overview.views.sheetView[0].showGridLines = True

    # Title Block
    ws_overview.merge_cells("A1:G1")
    title_cell = ws_overview["A1"]
    title_cell.value = f"NEXO PMO EXECUTIVE PROJECT OVERVIEW — {project.name} ({project.key})"
    title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    title_cell.fill = navy_fill
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_overview.row_dimensions[1].height = 36

    # Sub-header
    export_time = timezone.now().strftime("%Y-%m-%d %H:%M:%S UTC")
    exported_by_name = user.get_full_name() or user.username if user else "System Automated Export"
    ws_overview["A2"] = f"Generated: {export_time} | Exported By: {exported_by_name} | Platform: NEXO Enterprise PMO"
    ws_overview["A2"].font = Font(name="Segoe UI", size=9, italic=True, color="6B778C")
    ws_overview.row_dimensions[2].height = 20

    # Project Metadata Key-Value Table
    metadata_rows = [
        ("Project Key", project.key, "Lead / Owner", project.created_by.get_full_name() or project.created_by.username if project.created_by else "Unassigned"),
        ("Project Name", project.name, "Active Phase", current_phase),
        ("Health Status", health_status.replace("_", " "), "Total Tracker Tasks", issues_qs.count()),
        ("Completed Tasks", issues_qs.filter(status="DONE").count(), "In Progress Tasks", issues_qs.filter(status="IN_PROGRESS").count()),
        ("Open / To Do Tasks", issues_qs.filter(status="TODO").count(), "Critical Priority", issues_qs.filter(priority="CRITICAL").count()),
        ("Active Sprints", project.sprints.filter(status="ACTIVE").count(), "Attached Docs", project.docs.count()),
    ]

    r_idx = 4
    ws_overview.cell(row=r_idx, column=1, value="Core Project Telemetry").font = Font(name="Segoe UI", size=11, bold=True, color="0747A6")
    r_idx += 1

    for row_data in metadata_rows:
        ws_overview.cell(row=r_idx, column=1, value=row_data[0]).font = Font(name="Segoe UI", size=10, bold=True, color="172B4D")
        ws_overview.cell(row=r_idx, column=1).fill = _zebra_fill("EBECF0")
        ws_overview.cell(row=r_idx, column=1).border = border

        ws_overview.cell(row=r_idx, column=2, value=sanitize_formula_injection(row_data[1])).font = _cell_font()
        ws_overview.cell(row=r_idx, column=2).border = border

        ws_overview.cell(row=r_idx, column=3, value=row_data[2]).font = Font(name="Segoe UI", size=10, bold=True, color="172B4D")
        ws_overview.cell(row=r_idx, column=3).fill = _zebra_fill("EBECF0")
        ws_overview.cell(row=r_idx, column=3).border = border

        ws_overview.cell(row=r_idx, column=4, value=sanitize_formula_injection(row_data[3])).font = _cell_font()
        ws_overview.cell(row=r_idx, column=4).border = border
        ws_overview.row_dimensions[r_idx].height = 20
        r_idx += 1

    r_idx += 1
    # Health Rationale Section
    ws_overview.cell(row=r_idx, column=1, value="Executive Health Assessment & Rationale").font = Font(name="Segoe UI", size=11, bold=True, color="0747A6")
    r_idx += 1
    ws_overview.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx + 1, end_column=6)
    cell_hr = ws_overview.cell(row=r_idx, column=1, value=sanitize_formula_injection(health_rationale))
    cell_hr.font = _cell_font()
    cell_hr.alignment = Alignment(wrap_text=True, vertical="top")
    for r in range(r_idx, r_idx + 2):
        for c in range(1, 7):
            ws_overview.cell(row=r, column=c).border = border
    r_idx += 3

    # Executive Briefing Bullets
    ws_overview.cell(row=r_idx, column=1, value="Executive Briefing & Key Milestones").font = Font(name="Segoe UI", size=11, bold=True, color="0747A6")
    r_idx += 1
    bullets = [b.strip() for b in executive_bullets.split("\n") if b.strip()] if executive_bullets else ["No briefing notes available."]
    for bullet in bullets:
        ws_overview.cell(row=r_idx, column=1, value=sanitize_formula_injection(bullet)).font = _cell_font()
        ws_overview.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=6)
        ws_overview.row_dimensions[r_idx].height = 19
        r_idx += 1

    _auto_fit_columns(ws_overview, min_width=15, max_width=50)

    # =========================================================================
    # SHEET 2: TASKS
    # =========================================================================
    ws_tasks = wb.create_sheet(title="Tasks")
    ws_tasks.views.sheetView[0].showGridLines = True

    task_headers = [
        "Issue Key",
        "Workstream",
        "Title",
        "Type",
        "Status",
        "Priority",
        "Assignee Username",
        "Assignee Full Name",
        "Reporter",
        "Due Date",
        "Resolution",
        "Sprint",
        "Created At",
        "Updated At",
    ]

    ws_tasks.append(task_headers)
    ws_tasks.row_dimensions[1].height = 26
    for col_num in range(1, len(task_headers) + 1):
        cell = ws_tasks.cell(row=1, column=col_num)
        cell.font = _header_font()
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    t_idx = 2
    for issue in issues_qs:
        assignee_username = issue.assignee.username if issue.assignee else "Unassigned"
        assignee_fullname = issue.assignee.get_full_name() if issue.assignee else ""
        reporter_name = issue.reporter.username if issue.reporter else ""
        sprint_name = issue.sprint.name if issue.sprint else ""
        due_date_str = issue.due_date.strftime("%Y-%m-%d") if issue.due_date else ""
        created_str = issue.created_at.strftime("%Y-%m-%d %H:%M") if issue.created_at else ""
        updated_str = issue.updated_at.strftime("%Y-%m-%d %H:%M") if issue.updated_at else ""

        row = [
            f"{project.key}-{issue.id}",
            sanitize_formula_injection(issue.workstream or "General"),
            sanitize_formula_injection(issue.title),
            issue.issue_type,
            issue.status,
            issue.priority,
            sanitize_formula_injection(assignee_username),
            sanitize_formula_injection(assignee_fullname),
            sanitize_formula_injection(reporter_name),
            due_date_str,
            issue.resolution or ("Resolved" if issue.status == "DONE" else "Unresolved"),
            sanitize_formula_injection(sprint_name),
            created_str,
            updated_str,
        ]
        ws_tasks.append(row)
        ws_tasks.row_dimensions[t_idx].height = 20

        is_even = (t_idx % 2 == 0)
        for col_num in range(1, len(task_headers) + 1):
            c = ws_tasks.cell(row=t_idx, column=col_num)
            c.font = _cell_font()
            c.border = border
            if is_even:
                c.fill = zebra
        t_idx += 1

    _auto_fit_columns(ws_tasks, min_width=12, max_width=60)

    # =========================================================================
    # SHEET 3: MILESTONES
    # =========================================================================
    ws_milestones = wb.create_sheet(title="Milestones")
    ws_milestones.views.sheetView[0].showGridLines = True

    milestone_headers = [
        "Milestone / Deliverable",
        "Workstream / Phase",
        "Baseline Start",
        "Baseline Due",
        "Revised Due",
        "Schedule Variance",
        "Execution Status",
        "Schedule Status",
        "Owner / Assignee",
        "Dependencies & Notes",
    ]

    ws_milestones.append(milestone_headers)
    ws_milestones.row_dimensions[1].height = 26
    for col_num in range(1, len(milestone_headers) + 1):
        cell = ws_milestones.cell(row=1, column=col_num)
        cell.font = _header_font()
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    m_idx = 2
    ranked_deliverables = norm_state.get("ranked_deliverables", [])
    if not ranked_deliverables:
        # Fallback to MILESTONE signals if ranked_deliverables not populated
        milestone_signals = ProjectSignal.objects.filter(project=project, category="MILESTONE").order_by("target_date")
        for sig in milestone_signals:
            ws_milestones.append([
                sanitize_formula_injection(sig.title),
                sanitize_formula_injection(sig.description or "General"),
                "",
                sig.target_date.strftime("%Y-%m-%d") if sig.target_date else "",
                sig.target_date.strftime("%Y-%m-%d") if sig.target_date else "",
                "0 days",
                sig.status,
                "On Track",
                "Unassigned",
                sanitize_formula_injection(sig.source_location or ""),
            ])
            for col_num in range(1, len(milestone_headers) + 1):
                c = ws_milestones.cell(row=m_idx, column=col_num)
                c.font = _cell_font()
                c.border = border
            m_idx += 1
    else:
        for deliv in ranked_deliverables:
            ext_days = deliv.get("extension_days", 0)
            variance_str = f"+{ext_days} days" if ext_days and ext_days > 0 else "0 days"
            ws_milestones.append([
                sanitize_formula_injection(deliv.get("title", "")),
                sanitize_formula_injection(deliv.get("workstream") or deliv.get("phase") or "General"),
                deliv.get("baseline_start_date") or deliv.get("start_date") or "",
                deliv.get("baseline_end_date") or deliv.get("end_date") or "",
                deliv.get("end_date") or "",
                variance_str,
                deliv.get("execution_status") or deliv.get("status") or "PLANNED",
                deliv.get("schedule_status") or ("EXTENDED" if ext_days > 0 else "ON_TRACK"),
                sanitize_formula_injection(deliv.get("owner") or deliv.get("assignee") or "Unassigned"),
                sanitize_formula_injection(deliv.get("dependency") or ""),
            ])
            is_even = (m_idx % 2 == 0)
            for col_num in range(1, len(milestone_headers) + 1):
                c = ws_milestones.cell(row=m_idx, column=col_num)
                c.font = _cell_font()
                c.border = border
                if is_even:
                    c.fill = zebra
            m_idx += 1

    _auto_fit_columns(ws_milestones, min_width=14, max_width=50)

    # =========================================================================
    # SHEET 4: RISKS & ISSUES
    # =========================================================================
    ws_risks = wb.create_sheet(title="Risks & Issues")
    ws_risks.views.sheetView[0].showGridLines = True

    risk_headers = [
        "Category",
        "Title",
        "Severity",
        "Status",
        "Target Date",
        "Source Document",
        "Source Location",
        "Confidence",
        "Description",
    ]

    ws_risks.append(risk_headers)
    ws_risks.row_dimensions[1].height = 26
    for col_num in range(1, len(risk_headers) + 1):
        cell = ws_risks.cell(row=1, column=col_num)
        cell.font = _header_font()
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    r_idx = 2
    risk_signals = ProjectSignal.objects.filter(
        project=project, category__in=["RISK", "BLOCKER", "CHALLENGE", "DEPENDENCY"]
    ).select_related("source_doc").order_by("-confidence", "-id")

    for sig in risk_signals:
        doc_title = sig.source_doc.title if sig.source_doc else "Direct Input"
        t_date = sig.target_date.strftime("%Y-%m-%d") if sig.target_date else ""
        row = [
            sig.category,
            sanitize_formula_injection(sig.title),
            sig.severity,
            sig.status,
            t_date,
            sanitize_formula_injection(doc_title),
            sanitize_formula_injection(sig.source_location or ""),
            f"{int(sig.confidence * 100)}%",
            sanitize_formula_injection(sig.description or ""),
        ]
        ws_risks.append(row)
        is_even = (r_idx % 2 == 0)
        for col_num in range(1, len(risk_headers) + 1):
            c = ws_risks.cell(row=r_idx, column=col_num)
            c.font = _cell_font()
            c.border = border
            if is_even:
                c.fill = zebra
        r_idx += 1

    _auto_fit_columns(ws_risks, min_width=14, max_width=60)

    # =========================================================================
    # SHEET 5: LISTS & DEFINITIONS
    # =========================================================================
    ws_defs = wb.create_sheet(title="Lists & Definitions")
    ws_defs.views.sheetView[0].showGridLines = True

    ws_defs.merge_cells("A1:E1")
    head_def = ws_defs["A1"]
    head_def.value = "NEXO PMO REFERENCE DICTIONARY & ROSTER"
    head_def.font = _header_font(size=12)
    head_def.fill = navy_fill
    head_def.alignment = Alignment(horizontal="left", indent=1)
    ws_defs.row_dimensions[1].height = 28

    # Members Table
    ws_defs.cell(row=3, column=1, value="Project Team Member Roster").font = Font(name="Segoe UI", size=11, bold=True, color="0747A6")
    member_headers = ["Username", "Full Name", "Email Address", "Project Role", "Designation"]
    ws_defs.append(member_headers)
    ws_defs.row_dimensions[4].height = 22
    for c_i in range(1, 6):
        c = ws_defs.cell(row=4, column=c_i)
        c.font = _header_font(size=10)
        c.fill = sub_fill
        c.border = border

    cur_r = 5
    memberships = project.memberships.select_related("user").all()
    for m in memberships:
        u = m.user
        ws_defs.append([
            sanitize_formula_injection(u.username),
            sanitize_formula_injection(u.get_full_name()),
            sanitize_formula_injection(u.email),
            m.role,
            sanitize_formula_injection(getattr(u, "designation", "")),
        ])
        for c_i in range(1, 6):
            c = ws_defs.cell(row=cur_r, column=c_i)
            c.font = _cell_font()
            c.border = border
        cur_r += 1

    cur_r += 2
    # Standard Workflow Status Definitions Table
    ws_defs.cell(row=cur_r, column=1, value="Workflow States & Categories").font = Font(name="Segoe UI", size=11, bold=True, color="0747A6")
    cur_r += 1
    ws_defs.cell(row=cur_r, column=1, value="State Name").font = _header_font(size=10)
    ws_defs.cell(row=cur_r, column=1).fill = sub_fill
    ws_defs.cell(row=cur_r, column=1).border = border
    ws_defs.cell(row=cur_r, column=2, value="Category Mapping").font = _header_font(size=10)
    ws_defs.cell(row=cur_r, column=2).fill = sub_fill
    ws_defs.cell(row=cur_r, column=2).border = border
    ws_defs.cell(row=cur_r, column=3, value="Description").font = _header_font(size=10)
    ws_defs.cell(row=cur_r, column=3).fill = sub_fill
    ws_defs.cell(row=cur_r, column=3).border = border
    cur_r += 1

    statuses = [
        ("To Do", "TODO", "Work item identified, backlog item, or planned for upcoming iteration."),
        ("In Progress", "IN_PROGRESS", "Actively being developed, executed, or under active validation."),
        ("Done", "DONE", "Completed, verified against acceptance criteria, and deployed."),
    ]
    for st_name, cat, desc in statuses:
        ws_defs.cell(row=cur_r, column=1, value=st_name).font = _cell_font()
        ws_defs.cell(row=cur_r, column=1).border = border
        ws_defs.cell(row=cur_r, column=2, value=cat).font = _cell_font()
        ws_defs.cell(row=cur_r, column=2).border = border
        ws_defs.cell(row=cur_r, column=3, value=desc).font = _cell_font()
        ws_defs.cell(row=cur_r, column=3).border = border
        cur_r += 1

    cur_r += 2
    # Priority Levels Table
    ws_defs.cell(row=cur_r, column=1, value="Priority Levels").font = Font(name="Segoe UI", size=11, bold=True, color="0747A6")
    cur_r += 1
    ws_defs.cell(row=cur_r, column=1, value="Priority").font = _header_font(size=10)
    ws_defs.cell(row=cur_r, column=1).fill = sub_fill
    ws_defs.cell(row=cur_r, column=1).border = border
    ws_defs.cell(row=cur_r, column=2, value="SLA / Impact").font = _header_font(size=10)
    ws_defs.cell(row=cur_r, column=2).fill = sub_fill
    ws_defs.cell(row=cur_r, column=2).border = border
    cur_r += 1

    priorities = [
        ("CRITICAL", "Immediate blocker impacting production, launch timeline, or key milestone."),
        ("HIGH", "Important deliverable on critical path with active dependencies."),
        ("MEDIUM", "Standard backlog task or story with planned sprint allocation."),
        ("LOW", "Minor enhancement, optional improvement, or cosmetic update."),
    ]
    for p_name, sla in priorities:
        ws_defs.cell(row=cur_r, column=1, value=p_name).font = _cell_font()
        ws_defs.cell(row=cur_r, column=1).border = border
        ws_defs.cell(row=cur_r, column=2, value=sla).font = _cell_font()
        ws_defs.cell(row=cur_r, column=2).border = border
        cur_r += 1

    _auto_fit_columns(ws_defs, min_width=15, max_width=60)

    # Save to in-memory buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_canonical_project_csv(project: Project, issues_qs: Optional[Any] = None) -> str:
    """
    Generates a sanitized CSV string of project issues with formula injection protection.
    """
    if issues_qs is None:
        issues_qs = project.issues.all().select_related("assignee", "reporter", "sprint")

    output = io.StringIO()
    writer = csv.writer(output)

    headers = [
        "Key",
        "Workstream",
        "Title",
        "Type",
        "Status",
        "Priority",
        "Assignee",
        "Reporter",
        "Due Date",
        "Resolution",
        "Sprint",
        "Created At",
        "Updated At",
    ]
    writer.writerow(headers)

    for issue in issues_qs:
        assignee_name = issue.assignee.username if issue.assignee else "Unassigned"
        reporter_name = issue.reporter.username if issue.reporter else ""
        sprint_name = issue.sprint.name if issue.sprint else ""
        due_str = issue.due_date.strftime("%Y-%m-%d") if issue.due_date else ""
        created_str = issue.created_at.strftime("%Y-%m-%d %H:%M") if issue.created_at else ""
        updated_str = issue.updated_at.strftime("%Y-%m-%d %H:%M") if issue.updated_at else ""

        row = [
            f"{project.key}-{issue.id}",
            sanitize_formula_injection(issue.workstream or "General"),
            sanitize_formula_injection(issue.title),
            issue.issue_type,
            issue.status,
            issue.priority,
            sanitize_formula_injection(assignee_name),
            sanitize_formula_injection(reporter_name),
            due_str,
            issue.resolution or ("Resolved" if issue.status == "DONE" else "Unresolved"),
            sanitize_formula_injection(sprint_name),
            created_str,
            updated_str,
        ]
        writer.writerow(row)

    return output.getvalue()
