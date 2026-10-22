import io
from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from issues.models import Issue
from projects.models import Project, ProjectDoc, ProjectIntelligenceSummary, ProjectMembership, ProjectSignal
from projects.exports import generate_canonical_project_workbook, generate_canonical_project_csv, sanitize_formula_injection

User = get_user_model()


class TaskManagementEnhancementTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.admin_user = User.objects.create_user(
            username="admin_lead",
            email="admin@nexo.test",
            password="password123",
            user_type="ADMIN",
        )
        self.member_user = User.objects.create_user(
            username="alice_dev",
            first_name="Alice",
            last_name="Engineer",
            email="alice@nexo.test",
            password="password123",
            user_type="MEMBER",
        )
        self.project = Project.objects.create(
            name="Alpha Platform",
            key="ALPHA",
            created_by=self.admin_user,
        )
        ProjectMembership.objects.create(
            project=self.project,
            user=self.admin_user,
            role=ProjectMembership.Role.ADMIN,
        )
        ProjectMembership.objects.create(
            project=self.project,
            user=self.member_user,
            role=ProjectMembership.Role.MEMBER,
        )

        # Intelligence summary with sample ranked deliverables
        self.summary = ProjectIntelligenceSummary.objects.create(
            project=self.project,
            health_status=ProjectIntelligenceSummary.HealthStatus.ON_TRACK,
            health_rationale="Execution is progressing according to baseline schedule.",
            current_phase="Core Ingestion",
            executive_summary="- Baseline ingestion deployed.\n- Silver schema validation active.",
            normalized_state={
                "ranked_deliverables": [
                    {
                        "title": "AE Raw Ingestion",
                        "status": "COMPLETED",
                        "execution_status": "COMPLETED",
                        "schedule_status": "ON_TRACK",
                        "has_extension": False,
                        "extension_days": 0,
                        "workstream": "Pipeline",
                        "owner": "alice_dev",
                        "end_date": "2026-10-15",
                    },
                    {
                        "title": "Silver Validation Pipeline",
                        "status": "IN_PROGRESS",
                        "execution_status": "IN_PROGRESS",
                        "schedule_status": "EXTENDED",
                        "has_extension": True,
                        "extension_days": 6,
                        "workstream": "Analytics",
                        "owner": "Alice Engineer",
                        "end_date": "2026-10-22",
                    },
                    {
                        "title": "=SUM(1,2) Ingestion Benchmark",
                        "status": "PLANNED",
                        "execution_status": "PLANNED",
                        "schedule_status": "ON_TRACK",
                        "has_extension": False,
                        "extension_days": 0,
                        "workstream": "Benchmarking",
                        "owner": "Unknown Dev",
                        "end_date": "2026-11-01",
                    },
                ]
            }
        )

        # Risk signal
        self.signal_risk = ProjectSignal.objects.create(
            project=self.project,
            category=ProjectSignal.Category.RISK,
            title="Database Connection Pool Saturation",
            severity=ProjectSignal.Severity.HIGH,
            status="ACTIVE",
            confidence=0.95,
        )

    def test_sanitize_formula_injection(self):
        """Formula injection dangerous prefixes are safely escaped with single quote."""
        self.assertEqual(sanitize_formula_injection("=cmd|'/C calc'!A0"), "'=cmd|'/C calc'!A0")
        self.assertEqual(sanitize_formula_injection("+SUM(1,2)"), "'+SUM(1,2)")
        self.assertEqual(sanitize_formula_injection("-10%"), "'-10%")
        self.assertEqual(sanitize_formula_injection("@SUM(A1:A5)"), "'@SUM(A1:A5)")
        self.assertEqual(sanitize_formula_injection("Normal Title"), "Normal Title")
        self.assertEqual(sanitize_formula_injection(None), "")

    def test_canonical_export_workbook_generation(self):
        """Export generates valid 5-sheet openpyxl workbook with sanitized cells."""
        # Create an issue with potentially dangerous formula prefix
        Issue.objects.create(
            project=self.project,
            title="@MALICIOUS_TAG",
            workstream="Security",
            status="TODO",
            priority="HIGH",
            reporter=self.admin_user,
        )

        buf = generate_canonical_project_workbook(self.project, user=self.admin_user)
        self.assertGreater(len(buf.getvalue()), 0)

        wb = load_workbook(buf)
        expected_sheets = ["Project Overview", "Tasks", "Milestones", "Risks & Issues", "Lists & Definitions"]
        self.assertEqual(wb.sheetnames, expected_sheets)

        # Verify Tasks sheet formula escaping
        ws_tasks = wb["Tasks"]
        row_values = [ws_tasks.cell(row=r, column=3).value for r in range(2, ws_tasks.max_row + 1)]
        self.assertIn("'@MALICIOUS_TAG", row_values)

        # Verify Milestones sheet contains deliverable with =SUM escaped
        ws_milestones = wb["Milestones"]
        m_titles = [ws_milestones.cell(row=r, column=1).value for r in range(2, ws_milestones.max_row + 1)]
        self.assertIn("'=SUM(1,2) Ingestion Benchmark", m_titles)

    def test_canonical_export_csv_generation(self):
        """Export CSV escapes formula prefixes."""
        Issue.objects.create(
            project=self.project,
            title="=2+5",
            status="TODO",
            priority="LOW",
            reporter=self.admin_user,
        )
        csv_str = generate_canonical_project_csv(self.project)
        self.assertIn("Key,Workstream,Title", csv_str)
        self.assertIn("'=2+5", csv_str)

    def test_import_preview_endpoint(self):
        """import_preview endpoint returns matched users, suggested status, and reconciliation states."""
        self.client.force_authenticate(user=self.admin_user)
        response = self.client.get(f"/api/projects/{self.project.id}/import_preview/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        data = response.data

        self.assertEqual(data["total_deliverables"], 3)
        self.assertEqual(data["new_count"], 3)
        self.assertEqual(data["updated_count"], 0)

        # Verify owner fuzzy matching
        items = data["items"]
        ae_raw = next(i for i in items if "AE Raw" in i["title"])
        self.assertEqual(ae_raw["status"], "DONE")
        self.assertEqual(ae_raw["matched_user_id"], self.member_user.id)
        self.assertEqual(ae_raw["matched_username"], "alice_dev")

        silver = next(i for i in items if "Silver" in i["title"])
        self.assertEqual(silver["status"], "IN_PROGRESS")
        self.assertEqual(silver["priority"], "HIGH")
        self.assertEqual(silver["matched_user_id"], self.member_user.id)

    def test_confirm_import_creates_issues(self):
        """confirm_import atomically commits deliverables into persistent Issue records."""
        self.client.force_authenticate(user=self.admin_user)

        preview_resp = self.client.get(f"/api/projects/{self.project.id}/import_preview/")
        items = preview_resp.data["items"]

        confirm_resp = self.client.post(
            f"/api/projects/{self.project.id}/confirm_import/",
            {"selected_items": items},
            format="json",
        )
        self.assertEqual(confirm_resp.status_code, status.HTTP_200_OK)
        self.assertEqual(confirm_resp.data["created"], 3)

        # Verify issues in database
        self.assertEqual(Issue.objects.filter(project=self.project).count(), 3)
        ae_issue = Issue.objects.filter(project=self.project, title="AE Raw Ingestion").first()
        self.assertIsNotNone(ae_issue)
        self.assertEqual(ae_issue.status, "DONE")
        self.assertEqual(ae_issue.assignee, self.member_user)
        self.assertEqual(ae_issue.workstream, "Pipeline")

        # Second import preview shows 3 unchanged
        preview_2 = self.client.get(f"/api/projects/{self.project.id}/import_preview/")
        self.assertEqual(preview_2.data["new_count"], 0)
        self.assertEqual(preview_2.data["unchanged_count"], 3)

    def test_ai_assistant_query_and_execute_flow(self):
        """AI assistant answers questions and proposes actions that require explicit confirmation."""
        self.client.force_authenticate(user=self.admin_user)

        # Step 1: Read query
        resp_query = self.client.post(
            f"/api/projects/{self.project.id}/ai_assistant/",
            {"message": "Summarize project health and blockers"},
            format="json",
        )
        self.assertEqual(resp_query.status_code, status.HTTP_200_OK)
        self.assertIn("reply", resp_query.data)

        # Step 2: Action proposal
        resp_action = self.client.post(
            f"/api/projects/{self.project.id}/ai_assistant/",
            {"message": "Create task Deploy Redis Ingestion Cache with Critical priority"},
            format="json",
        )
        self.assertEqual(resp_action.status_code, status.HTTP_200_OK)
        self.assertEqual(resp_action.data["type"], "PROPOSAL")
        proposal = resp_action.data["proposal"]
        self.assertEqual(proposal["action"], "CREATE_ISSUE")

        # Verify nothing written to DB before confirmation
        self.assertFalse(Issue.objects.filter(title__icontains="Redis").exists())

        # Step 3: Explicit confirmation (ai_execute)
        resp_exec = self.client.post(
            f"/api/projects/{self.project.id}/ai_execute/",
            {"proposal": proposal},
            format="json",
        )
        self.assertEqual(resp_exec.status_code, status.HTTP_200_OK)
        self.assertTrue(resp_exec.data["success"])

        # Verify issue now exists
        new_issue = Issue.objects.filter(project=self.project, title__icontains="Redis").first()
        self.assertIsNotNone(new_issue)
        self.assertEqual(new_issue.reporter, self.admin_user)

    def test_ai_assistant_natural_language_parameter_extraction(self):
        """AI assistant extracts title, assignee, deadline, and priority from natural language."""
        self.client.force_authenticate(user=self.admin_user)

        # Natural language prompt matching user's exact scenario
        resp = self.client.post(
            f"/api/projects/{self.project.id}/ai_assistant/",
            {"message": "Create New Task Named Closure Document And Assign It To Alice And Make Deadline As Thursday [22/10/2026]"},
            format="json",
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["type"], "PROPOSAL")

        proposal = resp.data["proposal"]
        fields = proposal["fields"]
        self.assertEqual(fields["title"], "Closure Document")
        self.assertEqual(fields["assignee_name"], "alice_dev")
        self.assertEqual(fields["assignee_id"], self.member_user.id)
        self.assertEqual(fields["due_date"], "2026-10-22")
        self.assertEqual(fields["priority"], "MEDIUM")

        # Execute confirmed proposal
        exec_resp = self.client.post(
            f"/api/projects/{self.project.id}/ai_execute/",
            {"proposal": proposal},
            format="json",
        )
        self.assertEqual(exec_resp.status_code, status.HTTP_200_OK)

        # Verify database record
        issue = Issue.objects.get(id=exec_resp.data["issue_id"])
        self.assertEqual(issue.title, "Closure Document")
        self.assertEqual(issue.assignee, self.member_user)
        self.assertEqual(str(issue.due_date), "2026-10-22")

    def test_ai_assistant_tool_and_project_inquiries(self):
        """AI assistant answers questions about NEXO platform and team workload."""
        self.client.force_authenticate(user=self.admin_user)

        # Query 1: How to export
        resp_export = self.client.post(
            f"/api/projects/{self.project.id}/ai_assistant/",
            {"message": "How do I export to Excel or CSV?"},
            format="json",
        )
        self.assertEqual(resp_export.status_code, status.HTTP_200_OK)
        self.assertEqual(resp_export.data["type"], "ANSWER")
        self.assertIn("canonical PMO format", resp_export.data["reply"])

        # Query 2: Workload inquiry
        resp_workload = self.client.post(
            f"/api/projects/{self.project.id}/ai_assistant/",
            {"message": "What is Alice working on?"},
            format="json",
        )
        self.assertEqual(resp_workload.status_code, status.HTTP_200_OK)
        self.assertEqual(resp_workload.data["type"], "ANSWER")
        self.assertIn("alice_dev", resp_workload.data["reply"])

